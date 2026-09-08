from datetime import date, datetime
from email import encoders
from email.header import Header
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from html import escape
from tkinter import *
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText
from openpyxl import load_workbook
from urllib.parse import unquote, urlparse
import mimetypes
import os
import re
import sys
import time
import unicodedata

import boto3

# ========================
# CONFIGURACION DE BLOQUES
# ========================

BLOQUE_TAMANO = 50
ESPERA_ENTRE_BLOQUES = 60  # segundos

BENEFICIOS = [
    "1 TRASLADO EN AMBULANCIA GRATUITO AL AÑO EN CASO DE URGENCIA REAL",
    "ASESORIA MEDICA TELEFONICA Y POR VIDEOCONFERENCIA GRATUITA",
    "CONSULTA A DOMICILIO CON MEDICOS GENRALES DE NUESTRA RED A PRECIO PREFERENCIAL. ES INDISPENSABLE AGENDAR PREVIA CITA EN NUESTRO SERVICIO DE CONCIERGE",
    "COBERTURA INDIVIDUAL",
    "AMBULANCIA CON PRECIO PREFERENCIAL",
    "ASESORIA NUTRICIONAL TELEFONICA Y POR VIDEOCONFERENCIA",
    "ASESORIA EMOCIONAL TELEFONICA Y POR VIDEOCONFERENCIA",
    "RED DENTAL CON DESCUENTOS A NIVEL NACIONAL",
    "RED VISUAL CON DESCUENTOS A NIVEL NACIONAL",
    "DESCUENTOS EN ESTABLECIMENTOS COMERCIALES RED TDCONSENTIDO A NIVEL NACIONAL",
    "SERVICIO FUNERARIO. APLICA TIEMPO DE ESPERA DE 90 DIAS A PARTIR DE LA CONTRATACION DE LA MEMBRESIA A EXCEPCION DE MUERTE ACCIDENTAL LA CUAL APLICA DESDE EL DIA 1. INHUMACION O CREMACION. RECOLECCION DEL CUERPO, TRASLADO NIV NACIONAL, ARREGLO ESTETICO, SALA DE VELACION 30 PERSONAS EN CIRCULACION O DOMICILIO SIN COSTO ADICIONAL, SERVICIO DE TANATOLOGIA, ASESORIA JURIDICA TESTAMENTARIA VIA TEL. PARA INHUMACION SERV DE EMBALSAMADO, ATAUD METALICO, TRASLADO EN CARROZA.PARA CREMACION URNA BASICA Y SERV DE CREMACION",
    "CUPONES CON PROMOCIONES ESPECIALES PARA APP MOVIL.",
    "FARMACIA EN LINEA CON ENVIO A DOMICILIO",
    "SEGURO DE ACCIDENTES PERSONALES. MUERTE ACCIDENTAL $200,000 APLICA DE 12 A 70 AÑOS PERDIDA DE MIEMBROS POR ACCIDENTE ESCALA B $30,000 REEMBOLSO DE GASTOS FUNERARIOS POR ACCIDENTE $30,000 REEMBOLSO DE GASTOS MEDICOS POR ACCIDENTE $20,000 APLICAN DE 0 A 70 AÑOS. COBERTURAS A NIVEL NACIONAL. LAS MEMBRESIAS NO SON ACUMULABLES. POLIZA DE SEGURO 2510030074730",
    "PLATAFORMA CON EXPEDIENTE CLINICO ELECTRÓNICO.",
    "1 CHECK UP GRATIS PARA TITULAR. INCLUYE QS6, BIOMETRIA HEMATICA, EXAMEN GENERAL DE ORINA, INTERPRETACION DE RESULTADOS POR NUESTRO CESDI.",
    "1. CONSULTA PRESENCIAL O VIDEOCONSULTA CON MEDICINA GENERAL, NUTRICION, PSICOLOGIA, COORDINAR Y AGENDAR A TRAVES DE NUESTRO CONCIERGE. CONSULTA CON MEDICOS DE LA RED A PRECIO TABULADO.",
    "2 VIDEOCONSULTAS DE PEDIATRIA O GINECOLOGIA O MEDICINA INTERNA",
    "3 CUPONES 2 X 1 PARA CINE, SOLICITAR EL BENEFICIO DEBERA COMUNICARSE A CONCIERGE, SE VERIFICARA QUE EL USUARIO SE ENCUENTRE ACTIVO AL MOMENTO DE LA SOLICITUD Y DE SER EL CASO, SE PROPORCIONARA UN CODIGO 2X1. CONSULTA RESTRICCIONES.",
    "1 VIDEOCONSULTA O CONSULTA PRESENCIAL DE ESPECIALIDAD DENTRO DE NUESTRA RED MEDICA SIN COSTO. PARA UTILIZAR ESTE BENEFICIO, ES INDISPENSABLE COORDINAR Y AGENDAR LA CITA A TRAVES DE NUESTRO CONCIERGE MEDICO.",
    "PROTOCOLO DE TRIAGE MEDICO PARA DETERMINAR NIVEL DE URGENCIA, DIAGNOSTICO Y ATENCION",
    "VIDEOCONSULTAS ILIMITADAS CON MEDICINA GENERAL, NUTRICION Y PSICOLOGIA CON MEDICOS DEL CENTRO DE CONTACTO DE SALUD DIGITAL.",
]

MESES_ABREVIADOS = {
    1: "ene",
    2: "feb",
    3: "mar",
    4: "abr",
    5: "may",
    6: "jun",
    7: "jul",
    8: "ago",
    9: "sep",
    10: "oct",
    11: "nov",
    12: "dic",
}

# ========================
# VARIABLES GLOBALES
# ========================

ruta_excel = ""
ruta_html = ""


def cargar_env(ruta=".env"):
    # Busca el .env tanto en la carpeta actual como junto al script.
    rutas = [
        os.path.join(os.getcwd(), ruta),
        os.path.join(os.path.dirname(os.path.abspath(__file__)), ruta),
    ]

    for ruta_env in dict.fromkeys(rutas):
        if not os.path.exists(ruta_env):
            continue

        with open(ruta_env, "r", encoding="utf-8") as archivo:
            for linea in archivo:
                linea = linea.strip()
                if not linea or linea.startswith("#") or "=" not in linea:
                    continue

                clave, valor = linea.split("=", 1)
                clave = clave.strip()
                valor = valor.strip().strip('"').strip("'")
                os.environ.setdefault(clave, valor)


cargar_env()

AWS_REGION = os.getenv("AWS_REGION", "")
AWS_ACCESS_KEY = os.getenv("AWS_ACCESS_KEY", "")
AWS_SECRET_KEY = os.getenv("AWS_SECRET_KEY", "")
SES_SOURCE_EMAIL = os.getenv("SES_SOURCE_EMAIL", "")


def limpiar_valor(valor):
    # Convierte valores de Excel a texto limpio y consistente.
    if valor is None:
        return ""

    texto = str(valor).strip()
    if texto.startswith("'"):
        texto = texto[1:].strip()

    return texto


def limpiar_fecha(valor):
    # La vigencia debe mostrarse como fecha, sin hora.
    if valor is None:
        return ""

    if isinstance(valor, datetime):
        return formatear_fecha(valor.date())

    if isinstance(valor, date):
        return formatear_fecha(valor)

    texto = limpiar_valor(valor)
    match = re.match(r"^(\d{4}-\d{2}-\d{2})(?:\s+00:00:00(?:\.0)?)?$", texto)
    if match:
        try:
            return formatear_fecha(datetime.strptime(match.group(1), "%Y-%m-%d").date())
        except ValueError:
            return texto

    return texto


def formatear_fecha(fecha):
    return f"{fecha.day:02d}/{MESES_ABREVIADOS[fecha.month]}/{fecha.year}"


def normalizar_header(valor):
    # Facilita comparar encabezados aunque cambien acentos, espacios o mayusculas.
    texto = limpiar_valor(valor).lower().replace(" ", "")
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in texto if not unicodedata.combining(c))


def resolver_columna(headers, *aliases):
    # Permite aceptar varios nombres posibles para una misma columna.
    for alias in aliases:
        if alias in headers:
            return headers[alias]

    raise ValueError(f"Falta una columna en el Excel. Se esperaba alguna de: {', '.join(aliases)}")


def validar_columnas(headers):
    # Traduce nombres reales del Excel a las claves internas del programa.
    return {
        "producto": resolver_columna(headers, "producto"),
        "numerotarjeta": resolver_columna(headers, "numerotarjeta", "numerodetarjeta"),
        "nomcompleto": resolver_columna(headers, "nomcompleto", "nombrecompleto"),
        "vig": resolver_columna(headers, "vig", "vigencia"),
        "etiquetalogistica03": resolver_columna(headers, "etiquetalogistica03", "correo"),
        "simgfrente": headers.get("simgfrente"),
    }


def obtener_indice_columnas(ws):
    # Lee la primera fila del .xlsx y genera el mapa encabezado -> indice.
    headers = {}

    for i, celda in enumerate(ws[1], start=0):
        header = normalizar_header(celda.value)
        if header:
            headers[header] = i

    return validar_columnas(headers)


def obtener_indice_columnas_xls(header_row):
    # Hace el mismo mapeo de encabezados para archivos .xls.
    headers = {}

    for i, valor in enumerate(header_row):
        header = normalizar_header(valor)
        if header:
            headers[header] = i

    return validar_columnas(headers)


def leer_destinatarios(ruta):
    # Carga destinatarios desde Excel y los normaliza a un formato unico.
    extension = os.path.splitext(ruta)[1].lower()
    if extension == ".xls":
        return leer_destinatarios_xls(ruta)

    wb = load_workbook(ruta, data_only=True)
    ws = wb.active
    columnas = obtener_indice_columnas(ws)
    destinatarios = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        # Cada fila se convierte a un diccionario que luego usa la plantilla.
        producto = limpiar_valor(row[columnas["producto"]])
        numero_tarjeta = limpiar_valor(row[columnas["numerotarjeta"]])
        nombre = limpiar_valor(row[columnas["nomcompleto"]])
        vigencia = limpiar_fecha(row[columnas["vig"]])
        email = limpiar_valor(row[columnas["etiquetalogistica03"]])

        img_frente = ""
        if columnas.get("simgfrente") is not None and columnas["simgfrente"] < len(row):
            img_frente = limpiar_valor(row[columnas["simgfrente"]])

        # Solo se procesan filas con una direccion de correo utilizable.
        if email and "@" in email:
            destinatarios.append({
                "producto": producto,
                "numero_tarjeta": numero_tarjeta,
                "nombre": nombre,
                "vigencia": vigencia,
                "email": email,
                "img_frente": img_frente,
            })

    return destinatarios


def leer_destinatarios_xls(ruta):
    # Mantiene soporte para .xls con la misma salida que la lectura de .xlsx.
    import xlrd

    workbook = xlrd.open_workbook(ruta)
    sheet = workbook.sheet_by_index(0)
    columnas = obtener_indice_columnas_xls(sheet.row_values(0))
    destinatarios = []

    for row_idx in range(1, sheet.nrows):
        row = sheet.row_values(row_idx)
        # Se replica el mismo formato interno de destinatario.
        producto = limpiar_valor(row[columnas["producto"]])
        numero_tarjeta = limpiar_valor(row[columnas["numerotarjeta"]])
        nombre = limpiar_valor(row[columnas["nomcompleto"]])
        vigencia = limpiar_fecha(row[columnas["vig"]])
        email = limpiar_valor(row[columnas["etiquetalogistica03"]])

        img_frente = ""
        if columnas.get("simgfrente") is not None and columnas["simgfrente"] < len(row):
            img_frente = limpiar_valor(row[columnas["simgfrente"]])

        if email and "@" in email:
            destinatarios.append({
                "producto": producto,
                "numero_tarjeta": numero_tarjeta,
                "nombre": nombre,
                "vigencia": vigencia,
                "email": email,
                "img_frente": img_frente,
            })

    return destinatarios


def beneficios_html():
    # Inserta la lista fija de beneficios como HTML.
    items = "".join(f"<li>{escape(beneficio)}</li>" for beneficio in BENEFICIOS)
    return f'<ul style="margin:0; padding-left:18px;">{items}</ul>'


def imagen_frente_html(valor):
    # Si ya viene HTML lo respeta; si no, arma una etiqueta <img>.
    valor = limpiar_valor(valor)
    if not valor:
        return ""

    if valor.startswith("<"):
        return valor

    return f'<img src="{escape(valor, quote=True)}" alt="" style="max-width:180px; height:auto;">'


def personalizar_html(html_base, persona):
    # Reemplaza placeholders de la plantilla con los datos de cada persona.
    reemplazos = {
        "{$nomconcatenado}": escape(persona["nombre"]),
        "{$imgFrente}": imagen_frente_html(persona["img_frente"]),
        "{$membresia}": escape(persona["numero_tarjeta"]),
        "{$fVencimiento}": escape(persona["vigencia"]),
        "{$nomProducto}": escape(persona["producto"]),
        "{$nomBeneficio}": beneficios_html(),
        "{SimgFrente}": imagen_frente_html(persona["img_frente"]),
        "{Smembresia}": escape(persona["numero_tarjeta"]),
        "{SfVencimiento}": escape(persona["vigencia"]),
        "{SnomProducto}": escape(persona["producto"]),
        "{{NOMBRE}}": escape(persona["nombre"]),
        "{{EMPLEADO}}": escape(persona["numero_tarjeta"]),
        "{{VIGENCIA}}": escape(persona["vigencia"]),
        "{{PRODUCTO}}": escape(persona["producto"]),
    }

    html_personalizado = html_base
    for placeholder, valor in reemplazos.items():
        html_personalizado = html_personalizado.replace(placeholder, valor)

    return html_personalizado


def es_recurso_externo(src):
    # Evita intentar embeber URLs o esquemas que no son archivos locales.
    src = src.strip()
    esquema = urlparse(src).scheme.lower()
    return esquema in ("http", "https", "cid", "data", "mailto", "tel")


def ruta_imagen_local(carpeta_html, src):
    # Resuelve una imagen relativa al HTML y bloquea rutas fuera de esa carpeta.
    src_limpio = unquote(src.split("#", 1)[0].split("?", 1)[0]).replace("/", os.sep)
    ruta = os.path.abspath(os.path.join(carpeta_html, src_limpio))
    carpeta_base = os.path.abspath(carpeta_html)

    if not ruta.startswith(carpeta_base):
        return None

    if not os.path.isfile(ruta):
        return None

    return ruta


def embeber_imagenes_html(html, carpeta_html):
    # Convierte imagenes locales a cids para enviarlas embebidas por correo.
    imagenes = {}

    def reemplazar_src(match):
        inicio, src, fin = match.groups()
        if es_recurso_externo(src):
            return match.group(0)

        ruta = ruta_imagen_local(carpeta_html, src)
        if not ruta:
            return match.group(0)

        cid = os.path.basename(ruta)
        imagenes[cid] = ruta
        return f"{inicio}cid:{cid}{fin}"

    html_con_cid = re.sub(
        r'(<img\b[^>]*?\bsrc=["\'])([^"\']+)(["\'])',
        reemplazar_src,
        html,
        flags=re.IGNORECASE,
    )

    return html_con_cid, imagenes


def crear_adjunto_imagen(cid, ruta):
    # Prepara una imagen inline con el Content-ID que usa el HTML.
    content_type, _ = mimetypes.guess_type(ruta)
    if not content_type:
        content_type = "application/octet-stream"

    maintype, subtype = content_type.split("/", 1)

    with open(ruta, "rb") as f:
        adjunto = MIMEBase(maintype, subtype)
        adjunto.set_payload(f.read())

    encoders.encode_base64(adjunto)
    adjunto.add_header("Content-ID", f"<{cid}>")
    adjunto.add_header("Content-Disposition", "inline", filename=os.path.basename(ruta))

    return adjunto


def construir_mensaje_raw(asunto, remitente, destinatario, html_personalizado, carpeta_html):
    # Arma el mensaje MIME final con HTML e imagenes embebidas.
    html_con_cid, imagenes = embeber_imagenes_html(html_personalizado, carpeta_html)

    mensaje = MIMEMultipart("related")
    mensaje["Subject"] = Header(asunto, "utf-8")
    mensaje["From"] = remitente
    mensaje["To"] = destinatario

    alternative = MIMEMultipart("alternative")
    alternative.attach(MIMEText(html_con_cid, "html", "utf-8"))
    mensaje.attach(alternative)

    for cid, ruta in imagenes.items():
        mensaje.attach(crear_adjunto_imagen(cid, ruta))

    return mensaje


def crear_cliente_ses():
    # Valida configuracion minima y crea el cliente de AWS SES.
    faltantes = [
        nombre
        for nombre, valor in {
            "AWS_REGION": AWS_REGION,
            "AWS_ACCESS_KEY": AWS_ACCESS_KEY,
            "AWS_SECRET_KEY": AWS_SECRET_KEY,
        }.items()
        if not valor
    ]

    if faltantes:
        raise ValueError("Faltan variables en .env: " + ", ".join(faltantes))

    return boto3.client(
        "ses",
        region_name=AWS_REGION,
        aws_access_key_id=AWS_ACCESS_KEY,
        aws_secret_access_key=AWS_SECRET_KEY,
    )


def dominio_de_correo(correo):
    # Extrae el dominio del remitente para validar dominios verificados en SES.
    if "@" not in correo:
        return ""

    return correo.rsplit("@", 1)[1].strip().lower()


def obtener_identidades_verificadas(ses):
    # Recupera correos y dominios verificados para validar el remitente antes de enviar.
    identidades = []

    for tipo in ("EmailAddress", "Domain"):
        token = None
        while True:
            args = {
                "IdentityType": tipo,
                "MaxItems": 1000,
            }
            if token:
                args["NextToken"] = token

            response = ses.list_identities(**args)
            identidades.extend(response.get("Identities", []))
            token = response.get("NextToken")

            if not token:
                break

    verificadas = []
    for inicio in range(0, len(identidades), 100):
        # SES consulta atributos por bloques de hasta 100 identidades.
        bloque = identidades[inicio:inicio + 100]
        atributos = ses.get_identity_verification_attributes(
            Identities=bloque
        ).get("VerificationAttributes", {})

        for identidad in bloque:
            status = atributos.get(identidad, {}).get("VerificationStatus")
            if status == "Success":
                verificadas.append(identidad.lower())

    return verificadas


def validar_remitente_ses(ses, remitente):
    # El remitente es valido si el correo o su dominio estan verificados en SES.
    correo = remitente.strip().lower()
    dominio = dominio_de_correo(correo)
    verificadas = obtener_identidades_verificadas(ses)

    if correo in verificadas or dominio in verificadas:
        return

    opciones = "\n".join(f"- {identidad}" for identidad in verificadas[:20])
    if len(verificadas) > 20:
        opciones += f"\n- ... y {len(verificadas) - 20} mas"

    raise ValueError(
        f"El remitente no esta verificado en SES region {AWS_REGION}: {remitente}\n\n"
        "Usa uno de estos remitentes/dominios verificados:\n"
        f"{opciones or '- No se encontraron identidades verificadas'}"
    )


def escribir_log_ui(texto, tag):
    # Agrega una linea al log visual y mantiene visible el ultimo mensaje.
    log_text.configure(state=NORMAL)
    log_text.insert(END, texto + "\n", tag)
    log_text.see(END)
    log_text.configure(state=DISABLED)
    principal.update_idletasks()


def actualizar_ui(actual, total):
    # Sincroniza barra, contador y porcentaje durante el envio.
    progress["value"] = actual
    porcentaje = int((actual / total) * 100) if total else 0
    label_contador.config(text=f"{actual} / {total} correos")
    label_porcentaje.config(text=f"{porcentaje}%")
    principal.update_idletasks()


def limpiar_log_ui():
    # Reinicia el area de log antes de comenzar un nuevo proceso.
    log_text.configure(state=NORMAL)
    log_text.delete("1.0", END)
    log_text.configure(state=DISABLED)


# ========================
# INTERFAZ
# ========================

principal = Tk()
principal.title("Mailing Salud Interactiva")

if hasattr(sys, "_MEIPASS"):
    base_path = sys._MEIPASS
else:
    base_path = os.path.abspath(".")

icono = os.path.join(base_path, "si-logo.ico")
if os.path.exists(icono):
    principal.iconbitmap(icono)

principal.geometry("760x560")
principal.configure(bg="black")
principal.grid_columnconfigure(1, weight=1)

Label(principal, text="Remitente SES:", bg="black", fg="white").grid(row=0, column=0, padx=10, pady=5, sticky="e")
remitente_var = StringVar(value=SES_SOURCE_EMAIL)
Entry(principal, textvariable=remitente_var, width=70).grid(row=0, column=1, padx=5, pady=5, sticky="ew")

Label(principal, text="Asunto:", bg="black", fg="white").grid(row=1, column=0, padx=10, pady=5, sticky="e")
asunto_var = StringVar()
Entry(principal, textvariable=asunto_var, width=70).grid(row=1, column=1, padx=5, pady=5, sticky="ew")


# ========================
# SELECCION DE ARCHIVOS
# ========================

def seleccionar_excel():
    # Permite elegir el Excel con los destinatarios.
    global ruta_excel
    ruta_excel = filedialog.askopenfilename(
        title="Seleccionar Excel",
        filetypes=[("Excel files", "*.xlsx *.xlsm *.xls")],
    )
    if ruta_excel:
        label_excel.config(text=os.path.basename(ruta_excel))


def seleccionar_html():
    # Permite elegir la plantilla HTML del correo.
    global ruta_html
    ruta_html = filedialog.askopenfilename(
        title="Seleccionar HTML",
        filetypes=[("HTML files", "*.html *.htm")],
    )
    if ruta_html:
        label_html.config(text=os.path.basename(ruta_html))


# ========================
# FUNCION PRINCIPAL
# ========================

def enviar():
    # Coordina validaciones, lectura del Excel, personalizacion y envio de correos.
    global ruta_excel, ruta_html

    remitente = remitente_var.get().strip()
    asunto = asunto_var.get().strip()

    if not remitente or "@" not in remitente:
        messagebox.showerror("Error", "Escribe un correo remitente valido")
        return

    if not asunto:
        messagebox.showerror("Error", "Escribe el asunto del correo")
        return

    if not ruta_excel or not ruta_html:
        messagebox.showerror("Error", "Selecciona Excel y HTML")
        return

    carpeta_html = os.path.dirname(ruta_html)

    try:
        with open(ruta_html, "r", encoding="utf-8") as f:
            html_base = f.read()
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo leer el HTML:\n{e}")
        return

    try:
        destinatarios = leer_destinatarios(ruta_excel)
    except Exception as e:
        messagebox.showerror("Error", f"Excel invalido:\n{e}")
        return

    # Si no hay correos validos, no vale la pena continuar con SES.
    if not destinatarios:
        messagebox.showerror("Error", "No se encontraron correos en EtiquetaLogistica03")
        return

    try:
        ses = crear_cliente_ses()
        validar_remitente_ses(ses, remitente)
    except Exception as e:
        messagebox.showerror("AWS SES Error", str(e))
        return

    enviados = 0
    total = len(destinatarios)

    progress["maximum"] = total
    progress["value"] = 0
    actualizar_ui(0, total)
    limpiar_log_ui()

    nombre_log = f"log_envio_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"

    with open(nombre_log, "w", encoding="utf-8") as log_file:
        log_file.write("===== LOG DE ENVIO AWS SES =====\n")
        log_file.write(f"Fecha: {datetime.now()}\n")
        log_file.write(f"Remitente: {remitente}\n")
        log_file.write(f"Total destinatarios: {total}\n\n")

        for i, persona in enumerate(destinatarios, start=1):
            try:
                # Cada correo se construye de forma individual para personalizar su contenido.
                html_personalizado = personalizar_html(html_base, persona)
                mensaje = construir_mensaje_raw(
                    asunto,
                    remitente,
                    persona["email"],
                    html_personalizado,
                    carpeta_html,
                )

                response = ses.send_raw_email(
                    Source=remitente,
                    Destinations=[persona["email"]],
                    RawMessage={
                        "Data": mensaje.as_string(),
                    }
                )

                enviados += 1
                message_id = response.get("MessageId", "Sin MessageId")
                log_file.write(f"[OK] {persona['email']} - {message_id}\n")
                escribir_log_ui(f"[OK] {persona['email']} - {message_id}", "ok")

            except Exception as e:
                log_file.write(f"[ERROR] {persona['email']} -> {e}\n")
                escribir_log_ui(f"[ERROR] {persona['email']} -> {e}", "error")

            actualizar_ui(i, total)
            time.sleep(0.1)

            # Pausa entre bloques para evitar rafagas continuas de envios.
            if i % BLOQUE_TAMANO == 0 and i < total:
                mensaje_espera = f"Esperando {ESPERA_ENTRE_BLOQUES} segundos despues de {i} envios..."
                log_file.write(f"\n--- {mensaje_espera} ---\n\n")
                escribir_log_ui(mensaje_espera, "info")
                time.sleep(ESPERA_ENTRE_BLOQUES)

        log_file.write("\n===== RESUMEN =====\n")
        log_file.write(f"Enviados correctamente: {enviados}\n")
        log_file.write(f"No enviados: {total - enviados}\n")

    messagebox.showinfo(
        "Proceso terminado",
        f"{enviados} de {total} correos enviados.\nLog generado: {nombre_log}",
    )


# ========================
# BOTONES
# ========================

Button(principal, text="Seleccionar Excel", command=seleccionar_excel).grid(row=2, column=0, padx=10, pady=5, sticky="ew")

label_excel = Label(principal, text="Ningun archivo seleccionado", bg="black", fg="white", anchor="w")
label_excel.grid(row=2, column=1, padx=5, pady=5, sticky="ew")

Button(principal, text="Seleccionar HTML", command=seleccionar_html).grid(row=3, column=0, padx=10, pady=5, sticky="ew")

label_html = Label(principal, text="Ningun archivo seleccionado", bg="black", fg="white", anchor="w")
label_html.grid(row=3, column=1, padx=5, pady=5, sticky="ew")

progress = ttk.Progressbar(principal, orient="horizontal", length=500, mode="determinate")
progress.grid(row=4, column=1, padx=5, pady=15, sticky="ew")

label_contador = Label(principal, text="0 / 0 correos", bg="black", fg="white")
label_contador.grid(row=5, column=1, padx=5, sticky="w")

label_porcentaje = Label(principal, text="0%", bg="black", fg="white")
label_porcentaje.grid(row=5, column=1, padx=5, sticky="e")

Button(
    principal,
    text="Enviar Correos",
    command=enviar,
    bg="#215cc1",
    fg="white",
    width=30,
).grid(row=6, column=1, padx=5, pady=10)

Label(principal, text="Log en tiempo real:", bg="black", fg="white").grid(row=7, column=0, padx=10, pady=5, sticky="ne")

log_text = ScrolledText(principal, height=14, width=80, state=DISABLED, bg="#111111", fg="white")
log_text.grid(row=7, column=1, padx=5, pady=5, sticky="nsew")
log_text.tag_configure("ok", foreground="#33cc66")
log_text.tag_configure("error", foreground="#ff5555")
log_text.tag_configure("info", foreground="#d7d7d7")

principal.grid_rowconfigure(7, weight=1)

principal.mainloop()
